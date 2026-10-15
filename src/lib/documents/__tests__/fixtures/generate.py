"""Rebuild local synthetic fixtures; no downloaded or private documents."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from reportlab.pdfgen import canvas
from pypdf import PdfReader, PdfWriter
from docx import Document
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
import struct

ROOT = Path(__file__).parent
body = "A controlled trial found no significant difference in memory scores."
ref = "[1] Miller, G. (1956). The magical number seven. Psychological Review, 63(2), 81-97. doi:10.1037/h0043158"
pdf = canvas.Canvas(str(ROOT / "paper.pdf"))
pdf.setFont("Helvetica", 8)
for text in ["Synthetic evidence fixture", body, "References", ref]:
    pdf.drawString(40, 800 - ["Synthetic evidence fixture", body, "References", ref].index(text) * 35, text)
pdf.showPage()
pdf.drawString(40, 800, "Appendix")
pdf.drawString(40, 770, "Supplementary notes are outside the references section.")
pdf.save()
blank = canvas.Canvas(str(ROOT / "no-text.pdf"))
blank.rect(30, 30, 200, 200, fill=1)
blank.showPage()
blank.save()
encrypted = PdfWriter()
encrypted.append(PdfReader(ROOT / "paper.pdf"))
encrypted.encrypt("fixture-password")
encrypted.write(ROOT / "encrypted.pdf")
pages = PdfWriter()
for _ in range(101):
    pages.add_blank_page(width=100, height=100)
pages.write(ROOT / "too-many-pages.pdf")
document = Document()
document.add_heading("Synthetic evidence fixture", 0)
document.add_paragraph(body)
document.add_paragraph("中文正文：此文件仅用于安全导入测试。")
document.add_heading("参考文献", 1)
document.add_paragraph(ref)
document.add_heading("致谢", 1)
document.add_paragraph("This acknowledgement must not be selected as a reference.")
document.save(ROOT / "paper.docx")
(ROOT / "paper.txt").write_text(body + "\n参考文献\n" + ref + "\n致谢\nThanks.", encoding="utf-8")

def patch_docx(name, replacements):
    with ZipFile(ROOT / "paper.docx") as source, ZipFile(ROOT / name, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item.filename, replacements.pop(item.filename, source.read(item.filename)))
        for filename, data in replacements.items():
            target.writestr(filename, data)

patch_docx("external-reference.docx", {"word/_rels/document.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" TargetMode="External" Target="https://external.invalid/tracker"/></Relationships>'})
patch_docx("xml-entity.docx", {"word/document.xml": '<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///C:/private.txt">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&secret;</w:t></w:r></w:p></w:body></w:document>'})
patch_docx("zip-expansion.docx", {"word/document.xml": b"A" * (9 * 1024 * 1024)})
patch_docx("zip-path.docx", {"../escape.txt": "not extracted"})
(ROOT / "corrupt.pdf").write_bytes(b"%PDF-1.7\nthis is a corrupt synthetic PDF")
(ROOT / "corrupt.docx").write_bytes(b"PK\x03\x04corrupt")
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
chinese = canvas.Canvas(str(ROOT / "chinese.pdf"))
chinese.setFont("STSong-Light", 12)
chinese.drawString(40, 800, "中文研究正文：实验组与对照组没有显著差异。")
chinese.drawString(40, 760, "参考文献")
chinese.drawString(40, 720, "[1] 张三. 记忆研究. 2024.")
chinese.save()
patch_docx("xml-depth.docx", {"word/document.xml": '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + '<w:p>' * 150 + 'nested' + '</w:p>' * 150 + '</w:body></w:document>'})
with ZipFile(ROOT / "zip-size-lie.docx", "w", ZIP_DEFLATED) as target:
    target.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
    target.writestr("word/document.xml", "<text>" + "expanded " * 100_000 + "</text>")
raw = bytearray((ROOT / "zip-size-lie.docx").read_bytes())
offset = 0
while True:
    offset = raw.find(b"PK\x01\x02", offset)
    if offset < 0: break
    name_len = struct.unpack_from("<H", raw, offset + 28)[0]
    if raw[offset + 46:offset + 46 + name_len] == b"word/document.xml":
        struct.pack_into("<I", raw, offset + 24, 10)
    offset += 4
(ROOT / "zip-size-lie.docx").write_bytes(raw)
raw = bytearray((ROOT / "paper.docx").read_bytes())
offset = 0
while True:
    offset = raw.find(b"PK\x01\x02", offset)
    if offset < 0: break
    flags = struct.unpack_from("<H", raw, offset + 8)[0]
    struct.pack_into("<H", raw, offset + 8, flags | 1)
    offset += 4
(ROOT / "encrypted-flag.docx").write_bytes(raw)
print("Created synthetic PDF, DOCX, TXT and security fixtures in", ROOT)
