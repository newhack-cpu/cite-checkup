"""生成“引用体检 · 老师试用记录表”xlsx：一张表从上到下填，黄色格子要填、下拉选择，逐条判断合计和 SUS 得分自动计算。
按 A4 竖版分页：三、五、七、九、十一各起一页。用法：python -I build_trial_form.py <输出路径.xlsx>"""
import math
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

OUT = sys.argv[1]
FONT = "微软雅黑"
INK, SEAL, FAINT, LINK = "1D1A16", "9B2C24", "736A5E", "0563C1"
W = {"A": 6, "B": 48, "C": 15, "D": 33}
TOTAL_W = sum(W.values())
URL = "https://www.modelscope.cn/studios/kgkvybu111111/cite-checkup?mode=full"
SUS_DOI = "https://doi.org/10.1201/9781498710411-35"

wb = Workbook()
ws = wb.active
ws.title = "试用记录表"
ws.sheet_view.showGridLines = False
for col, width in W.items():
    ws.column_dimensions[col].width = width
wb.properties.title = "引用体检 · 老师试用记录表"
wb.properties.creator = "引用体检"

thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
F_INPUT = PatternFill("solid", start_color="FFF2CC")
F_HEAD = PatternFill("solid", start_color="F2F2F2")
F_SECTION = PatternFill("solid", start_color="F6E7E4")
F_AUTO = PatternFill("solid", start_color="E2EFDA")
F_ORG = PatternFill("solid", start_color="EDEDED")


def font(size=10.5, bold=False, color=INK, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color, underline=underline)


def units(text):
    return sum(2 if ord(ch) > 255 else 1.15 for ch in text)


def lines(text, width):
    return sum(max(1, math.ceil(units(part) / (width * 0.95))) for part in str(text).split("\n"))


def height(texts, minimum=21):
    """按最长的那格估算行高（微软雅黑 10.5 号约 16.5 磅一行）"""
    return max(minimum, max(lines(t, w) for t, w in texts) * 16.5 + 7)


def put(ref, value=None, f=None, fill=None, border=False, h="left", v="center"):
    cell = ws[ref]
    if value is not None:
        cell.value = value
    cell.font = f or font()
    if fill:
        cell.fill = fill
    if border:
        cell.border = BOX
    cell.alignment = Alignment(wrap_text=True, horizontal=h, vertical=v)
    return cell


def span(row, first, last, fill=None, border=False):
    ws.merge_cells(f"{first}{row}:{last}{row}")
    cols = "ABCD"
    for col in cols[cols.index(first):cols.index(last) + 1]:
        if fill:
            ws[f"{col}{row}"].fill = fill
        if border:
            ws[f"{col}{row}"].border = BOX


validations = {}


def choice(name, items, strict=False, prompt=None):
    if name not in validations:
        dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
        dv.showErrorMessage = True
        dv.errorStyle = "stop" if strict else "warning"
        dv.errorTitle = "请从下拉列表选择"
        dv.error = "点格子右边的小三角，从列表里选一项。"
        if prompt:
            dv.showInputMessage = True
            dv.promptTitle = name
            dv.prompt = prompt
        ws.add_data_validation(dv)
        validations[name] = dv
    return validations[name]


def number(name, kind="decimal", top=None):
    if name not in validations:
        dv = DataValidation(type=kind, operator="between" if top else "greaterThanOrEqual", formula1="0", formula2=str(top) if top else None, allow_blank=True)
        dv.showErrorMessage = True
        dv.errorStyle = "warning"
        dv.errorTitle = "请填数字"
        dv.error = "这一格填数字即可。"
        ws.add_data_validation(dv)
        validations[name] = dv
    return validations[name]


r = 1


def gap(size=10):
    global r
    ws.row_dimensions[r].height = size
    r += 1


def para(text, size=10.5, color=INK, bold=False, link=None, minimum=20):
    global r
    span(r, "A", "D")
    cell = put(f"A{r}", text, font(size, bold, color, "single" if link else None))
    if link:
        cell.hyperlink = link
    ws.row_dimensions[r].height = height([(text, TOTAL_W)], minimum)
    r += 1


def section(text, fill=F_SECTION, color=SEAL, new_page=False):
    global r
    if new_page:
        ws.row_breaks.append(Break(id=r - 1))
    gap()
    span(r, "A", "D", fill=fill)
    put(f"A{r}", text, font(12, True, color), fill=fill)
    ws.row_dimensions[r].height = 26
    r += 1


def header(labels, merge_cd=False, fill=F_HEAD):
    global r
    for col, text in zip("ABCD", labels):
        if text is not None:
            put(f"{col}{r}", text, font(10, True), fill=fill, border=True, h="center")
    if merge_cd:
        span(r, "C", "D", fill=fill, border=True)
    ws.row_dimensions[r].height = 22
    r += 1


def row(no, label, dv=None, merge_cd=False, required=True, answer_height=None, fill=F_INPUT, fmt=None):
    """一道题：A 序号，B 题目，C 填写（黄色），D 补充说明（白色，可不填）；merge_cd 时 C:D 合并成一个大格"""
    global r
    put(f"A{r}", no, font(10.5, color=FAINT), border=True, h="center")
    put(f"B{r}", label, border=True)
    put(f"C{r}", None, fill=fill if required else None, border=True, h="left" if merge_cd else "center", v="top" if merge_cd else "center")
    if fmt:
        ws[f"C{r}"].number_format = fmt
    if dv:
        dv.add(f"C{r}")
    if merge_cd:
        span(r, "C", "D", fill=fill if required else None, border=True)
    else:
        put(f"D{r}", None, border=True)
    ws.row_dimensions[r].height = max(answer_height or 0, height([(label, W["B"])]))
    r += 1
    return r - 1


# ---------- 标题和说明 ----------
ws.row_dimensions[r].height = 34
span(r, "A", "D")
put(f"A{r}", "引用体检 · 老师试用记录表", font(16, True))
r += 1
para("请用一份真实的参考文献列表试用约 20–30 分钟，从上到下填完，连同下载的核查报告一起交回。")
para("黄色格子要填，带小三角的格子点开可以下拉选择；白色格子可以不填。", color=FAINT)
para("试用入口：" + URL, color=LINK, link=URL)
para("魔搭创空间，不用登录，电脑、手机都能打开；第一次打开要等约半分钟。", color=FAINT)
para("匿名填写，只写编号和学科；卡住了再问组织者，并在第五部分记下卡在哪一步。", color=FAINT)
para("隐私：普通体检把参考文献条目发给 Crossref、OpenAlex 等公开数据库，中文期刊论文还会把题名发给 DeepSeek 的搜索接口（搜索引擎，不是大模型）；“AI 继续查”、截图识别和论断核查会把相应文字发给 DeepSeek 的大模型。学生未公开的论文请先征得本人同意。", color=FAINT)

# ---------- 一、基本信息 ----------
section("一、基本信息")
header(["序", "项目", "填写", "补充说明（可不填）"])
first_input = row(1, "编号（组织者给，如 T1）")
row(2, "学科 / 院系")
row(3, "身份", choice("身份", ["教师", "研究生", "本科生", "其他"]))
row(4, "平时核对学生论文的参考文献", choice("频率", ["经常", "偶尔", "几乎不"]))
row(5, "设备（右边可写浏览器，如 Edge、Chrome、微信里打开）", choice("设备", ["电脑", "手机", "平板"]))
row(6, "试用的列表来自", choice("来源", ["学生论文", "自己的论文", "其他"]))
row(7, "参考文献共几条", number("条数", "whole", 999))
row(8, "其中中文几条", number("条数", "whole", 999))
row(9, "其中英文几条", number("条数", "whole", 999))
row(10, "放进工具的方式", choice("方式", ["粘贴文字", "导入文件", "截图识别", "多种方式"]))

# ---------- 二、用时 ----------
section("二、用时")
para("先用平时的方法核对其中 5 条并计时，再把整份列表交给工具体检。", color=FAINT)
header(["序", "项目", "填写", "补充说明（可不填）"])
row(1, "人工核对 5 条用时（分钟）", number("分钟"))
row(2, "人工核对用的方法（如知网、万方、谷歌学术、问 AI，可写多个）", merge_cd=True)
row(3, "人工核对发现的问题（第几条、什么问题）", merge_cd=True, answer_height=46)
row(4, "工具检查整份列表用时（分钟，从粘贴到看完结果）", number("分钟"))
row(5, "第一次打开网页等了多久（秒）", number("秒", "whole", 3600))
row(6, "中途有没有以为网页打不开", choice("有无", ["有", "没有"]))

# ---------- 三、逐条判断 ----------
section("三、逐条判断：工具标出的问题对不对", new_page=True)
para("工具标为有误、相近候选、查无此文、已撤稿、关注声明的条目，每条填一行，序号填工具里的编号。待人工和已核实的不用逐条填，在表下填个数。", color=FAINT)
header(["序号", "工具结论", "你的判断", "说明（哪里写错了，或为什么是误报）"])
verdict = choice("工具结论", ["有误", "相近候选", "查无此文", "已撤稿", "关注声明"])
judge = choice("你的判断", ["确实有问题", "误报", "拿不准"], prompt="确实有问题：这条引用确实写错或不存在；误报：其实没问题；拿不准：自己也判断不了。")
item_no = number("序号", "whole", 999)
first_item = r
for _ in range(15):
    put(f"A{r}", None, fill=F_INPUT, border=True, h="center")
    put(f"B{r}", None, fill=F_INPUT, border=True, h="center")
    put(f"C{r}", None, fill=F_INPUT, border=True, h="center")
    put(f"D{r}", None, border=True)
    item_no.add(f"A{r}")
    verdict.add(f"B{r}")
    judge.add(f"C{r}")
    ws.row_dimensions[r].height = 21
    r += 1
last_item = r - 1


def rng(col):
    return f"{col}{first_item}:{col}{last_item}"


for label, formula in [
    ("工具标出的条数", f"=COUNTA({rng('B')})"),
    ("其中确实有问题", f'=COUNTIF({rng("C")},"确实有问题")'),
    ("其中误报", f'=COUNTIF({rng("C")},"误报")'),
    ("其中拿不准", f'=COUNTIF({rng("C")},"拿不准")'),
]:
    put(f"A{r}", None, border=True)
    put(f"B{r}", label, border=True, h="right")
    put(f"C{r}", formula, fill=F_AUTO, border=True, h="center")
    put(f"D{r}", "自动计算", font(10, color=FAINT), border=True)
    ws.row_dimensions[r].height = 21
    r += 1
for label in ["待人工的条数", "已核实的条数"]:
    put(f"A{r}", None, border=True)
    put(f"B{r}", label, border=True, h="right")
    put(f"C{r}", None, fill=F_INPUT, border=True, h="center")
    number("条数", "whole", 999).add(f"C{r}")
    put(f"D{r}", None, border=True)
    ws.row_dimensions[r].height = 21
    r += 1

# ---------- 四、漏报 ----------
section("四、漏报：有问题但工具没标出来")
para("你知道有问题、工具却没标出来的条目，尤其是判“已核实”但其实有错的；没有就空着。", color=FAINT)
header(["序号", "工具结论", "实际问题", None], merge_cd=True)
missed = choice("漏报的工具结论", ["已核实", "待人工", "相近候选", "有误", "其他"])
for _ in range(4):
    put(f"A{r}", None, fill=F_INPUT, border=True, h="center")
    put(f"B{r}", None, fill=F_INPUT, border=True, h="center")
    put(f"C{r}", None, fill=F_INPUT, border=True, v="top")
    span(r, "C", "D", fill=F_INPUT, border=True)
    item_no.add(f"A{r}")
    missed.add(f"B{r}")
    ws.row_dimensions[r].height = 24
    r += 1

# ---------- 五、使用过程 ----------
section("五、使用过程：哪一步卡住了", new_page=True)
para("每项选一个完成情况；问了人或没完成的，在右边写卡在哪、问了什么。", color=FAINT)
header(["序", "任务", "完成情况", "卡在哪 / 问了什么"])
done = choice("完成情况", ["自己完成", "问了人", "没完成", "没做"])
for i, task in enumerate([
    "看懂这个工具是做什么的",
    "把参考文献放进去并开始体检",
    "看懂每条的结论（印章和下面那句话）",
    "点“需要处理”，只看有问题的条目",
    "点“详情”，看哪里不一致和规范写法",
    "对一条“待人工”点“AI 继续查”",
    "复制修正版清单或规范参考文献，或下载核查报告",
    "（选做）在“正文论断核查”里查一句论文里的话",
], 1):
    row(i, task, done)
row(9, "新手引导（第一次打开时右下角的“带我看看”）；右边写有没有帮助", choice("新手引导", ["跟着走完", "走了一部分", "跳过", "没注意到"]))

# ---------- 六、理解检查 ----------
section("六、理解检查")
para("不用回页面看，按你的理解每题选一个字母。", color=FAINT)
header(["题", "题目与选项", "你的选择", "补充（可不填）"])
letter = choice("你的选择", ["A", "B", "C", "D"], strict=True)
for i, (stem, options) in enumerate([
    ("某条盖了“查无此文”，说明：", ["这条文献一定是编造的", "数据库里没找到同一篇，疑似不存在，要人工确认", "网络出错，没查完", "说不清"]),
    ("某条盖了“待人工”，说明：", ["这条文献有问题", "工具已经查到，只是要老师确认", "开放数据库覆盖不到（如图书、学位论文），工具判断不了，要人工核对", "说不清"]),
    ("某条盖了“已核实”，说明：", ["文献内容支持论文里引用它的那句话", "引文格式完全符合学校要求", "找到同一篇，题名、作者、年份、期刊等一致", "说不清"]),
    ("某条盖了“相近候选”，说明：", ["只找到相似文献，不能确定是不是同一篇，要人工看", "信息有误，照工具给的写法改就行", "这篇文献已撤稿", "说不清"]),
], 1):
    text = stem + "\n" + "\n".join(f"{k}. {o}" for k, o in zip("ABCD", options))
    row(i, text, letter)

# ---------- 七、评价 ----------
section("七、评价", new_page=True)
para("1 = 完全不同意，5 = 完全同意，下拉选一个数字。", color=FAINT)
header(["序", "说法", "打分（1–5）", "补充（可不填）"])
scale = choice("打分", ["1", "2", "3", "4", "5"], strict=True, prompt="1 = 完全不同意，5 = 完全同意")
for i, text in enumerate([
    "不用别人教，我也能很快上手",
    "每条的结论和理由我看得懂",
    "我信任它标出的问题，愿意拿去提醒学生",
    "它能明显节省我核对参考文献的时间",
    "我会推荐学生在交论文前先用它自查",
], 1):
    row(i, text, scale)

# ---------- 八、SUS ----------
section("八、可用性量表（SUS）")
para("试用结束后马上填，凭第一感觉，填完不要回头改。1 = 非常不同意，5 = 非常同意。", color=FAINT)
header(["序", "说法", "打分（1–5）", "补充（可不填）"])
sus_scale = choice("SUS 打分", ["1", "2", "3", "4", "5"], strict=True, prompt="1 = 非常不同意，5 = 非常同意")
sus_rows = [row(i, text, sus_scale) for i, text in enumerate([
    "我想我会经常使用这个工具",
    "我觉得这个工具没必要这么复杂",
    "我觉得这个工具容易使用",
    "我想我需要有人帮忙才能用好这个工具",
    "我觉得这个工具的各项功能结合得很好",
    "我觉得这个工具前后不够一致",
    "我想大多数人都能很快学会使用这个工具",
    "我觉得这个工具用起来很别扭",
    "使用这个工具时我很有把握",
    "我需要先学很多东西才能开始用这个工具",
], 1)]
para("量表来源：System Usability Scale（Brooke，1996），" + SUS_DOI, size=9.5, color=FAINT, link=SUS_DOI)

# ---------- 九、开放问题 ----------
section("九、开放问题", new_page=True)
para("想到什么写什么，一两句就行。", color=FAINT)
header(["序", "问题", "回答", None], merge_cd=True)
for i, text in enumerate([
    "如果指导学生写论文，你会不会让学生先用它自查？为什么？",
    "要让你愿意经常用，它还缺什么？",
    "最有用的功能",
    "最困惑或最不满意的地方",
    "界面看着顺眼吗？哪里别扭？",
    "其他意见",
], 1):
    row(i, text, merge_cd=True, answer_height=46)

# ---------- 十、使用同意 ----------
section("十、使用同意")
para("每项选“同意”或“不同意”；两项都不同意时，本表只用于改进工具，不写进比赛材料。", color=FAINT)
header(["序", "用途", "选择", "补充（可不填）"])
agree = choice("同意", ["同意", "不同意"])
row(1, "在比赛材料中匿名引用我的意见（只写编号和学科）", agree)
row(2, "组织者用我交回的核查报告分析误报（报告只用于改进工具，不公开）", agree)
row(3, "愿意接受回访的联系方式（可不填）", merge_cd=True, required=False)
row(4, "日期", fmt="yyyy-mm-dd")

# ---------- 十一、组织者记录 ----------
section("十一、组织者记录（老师不用填）", fill=F_ORG, color=FAINT, new_page=True)
header(["序", "项目", "记录", None], merge_cd=True, fill=F_ORG)
row(1, "试用日期", merge_cd=True, fill=F_ORG)
row(2, "入口（魔搭 / 本机 / 其他）", merge_cd=True, fill=F_ORG)
row(3, "组织者协助几次、在哪一步", merge_cd=True, fill=F_ORG)
row(4, "出现的限流或出错提示（照抄原文）", merge_cd=True, fill=F_ORG)
row(5, "理解检查答对几题（满分 4，答案见试用指南第四节）", merge_cd=True, fill=F_ORG)
s = [f"C{n}" for n in sus_rows]
sus_formula = (
    f'=IF(COUNTA(C{sus_rows[0]}:C{sus_rows[-1]})<10,"未填完",'
    f"(({s[0]}-1)+(5-{s[1]})+({s[2]}-1)+(5-{s[3]})+({s[4]}-1)+(5-{s[5]})+({s[6]}-1)+(5-{s[7]})+({s[8]}-1)+(5-{s[9]}))*2.5)"
)
sus_row = row(6, "SUS 得分（0–100，自动计算：单数题减 1，双数题用 5 减，十题相加乘 2.5）", merge_cd=True, fill=F_AUTO)
ws[f"C{sus_row}"].value = sus_formula
ws[f"C{sus_row}"].alignment = Alignment(horizontal="left", vertical="center")
row(7, "交回的报告文件名", merge_cd=True, fill=F_ORG)
last_row = r - 1

# ---------- 打印与打开时的位置 ----------
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = "portrait"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.55, bottom=0.6, header=0.3, footer=0.3)
ws.print_options.horizontalCentered = True
ws.oddFooter.center.text = "引用体检 · 老师试用记录表　第 &P 页，共 &N 页"
ws.oddFooter.center.font = FONT
ws.oddFooter.center.size = 9
ws.print_area = f"A1:D{last_row}"
ws.sheet_view.selection[0].activeCell = f"C{first_input}"
ws.sheet_view.selection[0].sqref = f"C{first_input}"

wb.save(OUT)
print({"rows": last_row, "first_item": first_item, "last_item": last_item, "sus_rows": sus_rows, "sus_row": sus_row})
