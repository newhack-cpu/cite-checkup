"""Create public demo documents and measure local Mandarin narration; no credentials."""
from pathlib import Path
import json, subprocess
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import simpleSplit
from docx import Document

ROOT = Path(__file__).resolve().parents[3]
MEDIA = Path(__file__).resolve().parent
OUT = ROOT / 'eval/phase-two-2026-10-14/demo-inputs'
OUT.mkdir(parents=True, exist_ok=True)
refs = [
    '[1] LECUN Y, BENGIO Y, HINTON G | Deep learning | Nature | 2015 | DOI: 10.1038/nature14539',
    '[2] Hochreiter, S., & Schmidhuber, J. (1997). Long short-term memory. Neural Networks, 9(8), 1735-1780. DOI: 10.1162/neco.1997.9.8.1735',
]
paragraphs = ['Cite Checkup - Public Reference Demo', 'This demonstration contains two public references. The second intentionally has an incorrect journal name to demonstrate field verification.', 'References', *refs]
(OUT / 'public-references.txt').write_text('\n\n'.join(paragraphs), encoding='utf-8')
doc = Document()
for i, line in enumerate(paragraphs):
    doc.add_heading(line, 0 if i == 0 else 1) if i in (0, 2) else doc.add_paragraph(line)
doc.save(OUT / 'public-references.docx')
pdf = canvas.Canvas(str(OUT / 'public-references.pdf'), pagesize=A4)
pdf.setTitle('Cite Checkup - Public Reference Demo')
y = 790
for i, para in enumerate(paragraphs):
    font, size = ('Helvetica-Bold', 16) if i == 0 else ('Helvetica', 11)
    pdf.setFont(font, size)
    for line in simpleSplit(para, font, size, 475):
        pdf.drawString(60, y, line)
        y -= size * 1.7
    y -= 15
pdf.save()
timings = []
for wav in sorted((MEDIA / 'narration').glob('*.wav')):
    seconds = float(subprocess.check_output(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'default=noprint_wrappers=1:nokey=1', str(wav)], text=True))
    timings.append({'file': wav.name, 'seconds': seconds})
(MEDIA / 'narration/durations.json').write_text(json.dumps(timings, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'demoFiles': str(OUT), 'narrationSeconds': sum(x['seconds'] for x in timings)}, ensure_ascii=False))
